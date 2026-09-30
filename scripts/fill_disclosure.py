"""Fill the organisers' AI Usage Disclosure form (LangAI3.0_AI_Disclosure.docx) -> submission/.

    python scripts/fill_disclosure.py --template "<path to LangAI3.0_AI_Disclosure.docx>"
Signature is intentionally left blank for the team representative to sign.
"""
from __future__ import annotations

import argparse
import copy
from pathlib import Path

import docx

ROOT = Path(__file__).resolve().parents[1]

FIELDS = {
    "Team Name:": "TapTrace (Thapar Institute of Engineering and Technology)",
    "Project / Product Name:": "TapTrace - Smart Guided Troubleshooting Engine (Theme 02)",
    "Organization / Institution (if any):": "Thapar Institute of Engineering and Technology, Patiala",
    "Submission Date:": "30 September 2026",
    "Idea generation / brainstorming": "Yes - ideas were discussed together with Claude Code (Anthropic); final choices made by the team member.",
    "Code generation or assistance": "Yes - code was written together with Claude Code; equal contribution with the team member.",
    "UI / UX design": "Yes - demo page built together with Claude Code.",
    "Content creation": "Yes - README, docs and slides drafted together with Claude Code and reviewed by the team member.",
    "Data analysis": "Yes - analysis of the provided data files done together with Claude Code.",
    "Testing / debugging": "Yes - tests and debugging done together with Claude Code.",
    "Other": "The product calls Groq openai/gpt-oss-120b at runtime (query enrichment and plan structuring).",
}

FEATURES = [
    ("Grounded step extraction and relevance gate", "Both",
     "Claude Code (Anthropic), used together with the team member (equal contribution). Output: taptrace/siis.py. "
     "Reviewed and corrected by the team member against the 20 provided scenarios."),
    ("Deeplink catalog compiler and resolver", "Both",
     "Claude Code, used together with the team member (equal contribution). Output: taptrace/catalog.py, resolver.py; "
     "checked with a 37-case benchmark."),
    ("Semantic cache", "Both",
     "Claude Code, used together with the team member (equal contribution). Output: taptrace/cache.py."),
    ("API, Docker, demo, tests, evaluation, documentation and slides", "Both",
     "Claude Code, used together with the team member (equal contribution). Results are produced by the evaluation scripts."),
]


def set_after_label(p, value):
    runs = p.runs
    label_done = False
    for r in runs:
        if "_" in r.text and not label_done:
            r.text = " " + value
            label_done = True
        elif label_done and set(r.text.strip()) <= {"_"}:
            r.text = ""
    if not label_done:
        p.add_run(" " + value)


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--template", required=True)
    args = ap.parse_args()
    d = docx.Document(args.template)
    P = d.paragraphs
    for p in P:
        t = p.text.strip()
        for k, v in FIELDS.items():
            if t.startswith(k) and "_" in t:
                set_after_label(p, v)
        if t.startswith("Did your team use any Artificial Intelligence"):
            for r in p.runs:
                if "Yes / No" in r.text:
                    r.text = r.text.replace("Yes / No", "Yes")
    # the free line under the Yes/No question
    for i, p in enumerate(P):
        if p.text.strip().startswith("Did your team use"):
            nxt = P[i + 1]
            if set(nxt.text.strip()) <= {"_"}:
                for r in nxt.runs:
                    r.text = ""
                nxt.add_run("Yes. Claude Code (Anthropic), used together with the team member as an equal contribution.")
            break
    # feature blocks: template has 2 blocks of 3 paragraphs; fill first, clone as needed
    idx = [i for i, p in enumerate(P) if p.text.strip().startswith("1. Feature Name")]
    blocks = [[P[i], P[i + 1], P[i + 2]] for i in idx]
    anchor = blocks[-1][2]._p
    proto = [copy.deepcopy(x._p) for x in blocks[0]]
    while len(blocks) < len(FEATURES):
        new = [copy.deepcopy(x) for x in proto]
        for el in new:
            anchor.addnext(el)
            anchor = el
        blocks.append([docx.text.paragraph.Paragraph(el, blocks[0][0]._parent) for el in new])
    for (name, origin, desc), (p1, p2, p3) in zip(FEATURES, blocks):
        for p, v in ((p1, name), (p2, origin), (p3, desc)):
            txt = p.text
            for r in p.runs:
                r.text = ""
            head = txt.split("_")[0].rstrip()
            p.runs[0].text = head + " " + v if p.runs else None
    for p in P:
        t = p.text.strip()
        if t == "Name of Team Representative:":
            p.add_run(" Arnav Joshi")
        elif t == "Role:":
            p.add_run(" Team lead (solo participant)")
        elif t == "Date:":
            p.add_run(" 30 September 2026")
    out = ROOT / "submission" / "TIET_TapTrace_AI_Disclosure.docx"
    out.parent.mkdir(exist_ok=True)
    d.save(out)
    print("saved", out)


if __name__ == "__main__":
    main()
